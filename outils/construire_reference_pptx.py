"""Construit le fichier de référence PowerPoint du format cerema-pptx.

Le modèle « Diaporama Cerema 2026 » (.potx) n'est pas utilisable tel quel par
Pandoc, pour trois raisons :

- c'est un modèle (.potx), alors que Pandoc attend une présentation (.pptx) ;
- Pandoc choisit les dispositions par leur nom anglais (« Title Slide »,
  « Title and Content »…), alors que celles du modèle portent des noms français ;
- les titres du modèle sont des zones de texte génériques (type « body ») :
  Pandoc ne place un titre que dans une zone de type « title » ou « ctrTitle »,
  et le contenu dans une zone sans type (« obj »).

Le script retire les diapositives d'exemple, renomme sept dispositions et
retype leurs zones réservées. Les 33 autres dispositions restent disponibles
dans PowerPoint pour une mise en forme manuelle.

Usage :
    uv run python outils/construire_reference_pptx.py \
        sources/Modele_Diaporama_Cerema_2026.potx \
        _extensions/cerema/pptx/reference.pptx
"""

import copy
import io
import sys
import zipfile
from collections.abc import Callable
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.slide import SlideLayout

TYPE_MODELE = b"presentationml.template.main+xml"
TYPE_PRESENTATION = b"presentationml.presentation.main+xml"

# Géométrie commune (EMU) : diapositive 13 444 538 x 7 562 850.
GAUCHE = 1190610
"""Bord gauche du titre et du contenu dans le modèle."""
LARGEUR_UTILE = 11104000
"""Largeur de contenu, marges du modèle conservées à droite."""
HAUT_CONTENU = 1519867
"""Haut de la zone de contenu, sous le titre."""
BAS_CONTENU = 6700000
"""Bas de la zone de contenu, au-dessus du pied de page."""
GOUTTIERE = 457200
"""Écart entre deux colonnes (0,5 pouce)."""


def lire_modele(chemin: Path) -> Presentation:
    """Ouvre un modèle .potx comme une présentation .pptx.

    Args:
        chemin: Chemin du fichier .potx (ou .pptx).

    Returns:
        La présentation python-pptx, le type de contenu principal corrigé.
    """
    tampon = io.BytesIO()
    with (
        zipfile.ZipFile(chemin) as source,
        zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as cible,
    ):
        for entree in source.infolist():
            donnees = source.read(entree.filename)
            if entree.filename == "[Content_Types].xml":
                donnees = donnees.replace(TYPE_MODELE, TYPE_PRESENTATION)
            cible.writestr(entree, donnees)
    tampon.seek(0)
    return Presentation(tampon)


def supprimer_diapositives(presentation: Presentation) -> None:
    """Retire les diapositives d'exemple ; leurs images partent avec elles.

    Args:
        presentation: Présentation modifiée sur place.
    """
    liste = presentation.slides._sldIdLst
    for identifiant in list(liste):
        presentation.part.drop_rel(identifiant.get(qn("r:id")))
        liste.remove(identifiant)


def disposition(presentation: Presentation, nom: str) -> SlideLayout:
    """Cherche une disposition par son nom, espaces de bord ignorés.

    Args:
        presentation: Présentation où chercher.
        nom: Nom français de la disposition dans le modèle.

    Returns:
        La disposition trouvée.

    Raises:
        KeyError: Si aucune disposition ne porte ce nom.
    """
    for candidate in presentation.slide_layouts:
        if candidate.name.strip() == nom:
            return candidate
    raise KeyError(nom)


def zone(layout: SlideLayout, idx: int) -> etree._Element:
    """Renvoie l'élément <p:sp> de la zone réservée d'index ``idx``.

    Args:
        layout: Disposition à parcourir.
        idx: Index de la zone dans le modèle.

    Returns:
        L'élément XML de la forme.
    """
    for forme in layout.placeholders:
        if forme.placeholder_format.idx == idx:
            return forme._element
    raise KeyError(f"{layout.name} : zone {idx} absente")


def retyper(
    forme: etree._Element,
    type_zone: str | None,
    idx: int | None,
    sz: str | None = None,
) -> None:
    """Change le type et l'index d'une zone réservée.

    Args:
        forme: Élément <p:sp> de la zone.
        type_zone: Nouveau type (« title », « body »…) ; None pour « obj ».
        idx: Nouvel index ; None pour l'omettre (zones titre).
        sz: Taille indicative (« half », « quarter ») ou None.
    """
    ph = forme.find(".//" + qn("p:ph"))
    for attribut in ("type", "idx", "sz", "hasCustomPrompt"):
        ph.attrib.pop(attribut, None)
    if type_zone is not None:
        ph.set("type", type_zone)
    if sz is not None:
        ph.set("sz", sz)
    if idx is not None:
        ph.set("idx", str(idx))


def placer(forme: etree._Element, x: int, y: int, largeur: int, hauteur: int) -> None:
    """Fixe la position et la taille d'une zone (EMU).

    Args:
        forme: Élément <p:sp> de la zone.
        x: Abscisse du coin haut gauche.
        y: Ordonnée du coin haut gauche.
        largeur: Largeur.
        hauteur: Hauteur.
    """
    xfrm = forme.find(qn("p:spPr") + "/" + qn("a:xfrm"))
    xfrm.find(qn("a:off")).attrib.update({"x": str(x), "y": str(y)})
    xfrm.find(qn("a:ext")).attrib.update({"cx": str(largeur), "cy": str(hauteur)})


def titre_en_capitales(forme: etree._Element) -> None:
    """Impose les capitales au premier niveau du titre, comme le livre de marque.

    Args:
        forme: Élément <p:sp> de la zone de titre.
    """
    for niveau in forme.iter(qn("a:lvl1pPr")):
        rpr = niveau.find(qn("a:defRPr"))
        if rpr is not None:
            rpr.set("cap", "all")


def titre_diapositive(layout: SlideLayout) -> None:
    """Retype le titre d'une diapositive de contenu (zone 10) et l'élargit.

    Args:
        layout: Disposition modifiée sur place.
    """
    forme = zone(layout, 10)
    retyper(forme, "title", None)
    titre_en_capitales(forme)
    xfrm = forme.find(qn("p:spPr") + "/" + qn("a:xfrm"))
    xfrm.find(qn("a:ext")).set("cx", str(LARGEUR_UTILE))


def deux_colonnes() -> tuple[int, int]:
    """Calcule la largeur d'une colonne et l'abscisse de la colonne de droite.

    Returns:
        La largeur d'une colonne et l'abscisse de la colonne de droite (EMU).
    """
    largeur = (LARGEUR_UTILE - GOUTTIERE) // 2
    return largeur, GAUCHE + largeur + GOUTTIERE


def cloner(forme: etree._Element) -> etree._Element:
    """Duplique une zone juste après elle, avec un identifiant de forme neuf.

    Args:
        forme: Élément <p:sp> à dupliquer.

    Returns:
        La copie, insérée dans l'arbre.
    """
    copie = copy.deepcopy(forme)
    arbre = forme.getparent()
    ids = [int(e.get("id")) for e in arbre.iter(qn("p:cNvPr"))]
    cnv = copie.find(".//" + qn("p:cNvPr"))
    cnv.set("id", str(max(ids) + 1))
    for ext in cnv.findall(qn("a:extLst")):
        cnv.remove(ext)
    forme.addnext(copie)
    return copie


def vers_title_slide(layout: SlideLayout) -> None:
    """« Slide couv » : titre (12), sous-titre (10), date (11)."""
    titre = zone(layout, 12)
    largeur_titre = titre.find(
        qn("p:spPr") + "/" + qn("a:xfrm") + "/" + qn("a:ext")
    ).get("cx")
    retyper(titre, "ctrTitle", None)
    sous_titre = zone(layout, 10)
    retyper(sous_titre, "subTitle", 1)
    # Le sous-titre (et l'auteur, que Pandoc y ajoute) prend la largeur du titre.
    sous_titre.find(qn("p:spPr") + "/" + qn("a:xfrm") + "/" + qn("a:ext")).set(
        "cx", largeur_titre
    )
    date = zone(layout, 11)
    retyper(date, "dt", 10, "quarter")
    # Une zone « dt » hérite du style des autres textes du masque (Times) :
    # on fixe Marianne, la police du texte courant.
    for rpr in date.iter(qn("a:defRPr")):
        police = etree.SubElement(rpr, qn("a:latin"))
        police.set("typeface", "Marianne")


def vers_title_and_content(layout: SlideLayout) -> None:
    """« Slide texte seulement » : titre (10), contenu pleine largeur (11)."""
    titre_diapositive(layout)
    contenu = zone(layout, 11)
    retyper(contenu, None, 1)
    placer(contenu, GAUCHE, HAUT_CONTENU, LARGEUR_UTILE, BAS_CONTENU - HAUT_CONTENU)


def vers_section_header(layout: SlideLayout) -> None:
    """« Slide fond dégradé » : titre de chapitre (10)."""
    retyper(zone(layout, 10), "title", None)


def vers_two_content(layout: SlideLayout) -> None:
    """« Slide texte 2 colonnes » : titre (10), deux contenus côte à côte."""
    titre_diapositive(layout)
    gauche = zone(layout, 11)
    # La zone d'origine répartit son texte sur deux colonnes internes (numCol) :
    # chaque moitié doit n'en avoir qu'une.
    gauche.find(".//" + qn("a:bodyPr")).attrib.pop("numCol", None)
    droite = cloner(gauche)
    hauteur = BAS_CONTENU - HAUT_CONTENU
    largeur, x_droite = deux_colonnes()
    retyper(gauche, None, 1, "half")
    retyper(droite, None, 2, "half")
    placer(gauche, GAUCHE, HAUT_CONTENU, largeur, hauteur)
    placer(droite, x_droite, HAUT_CONTENU, largeur, hauteur)


def vers_comparison(layout: SlideLayout) -> None:
    """« Slide texte + graphiques » : deux intertitres, deux contenus dessous."""
    titre_diapositive(layout)
    hauteur_entete = 500000
    haut = HAUT_CONTENU + hauteur_entete + 100000
    hauteur = BAS_CONTENU - haut
    entete_gauche = zone(layout, 11)
    entete_droite = cloner(entete_gauche)
    largeur, x_droite = deux_colonnes()
    retyper(entete_gauche, "body", 1)
    retyper(entete_droite, "body", 3, "quarter")
    placer(entete_gauche, GAUCHE, HAUT_CONTENU, largeur, hauteur_entete)
    placer(entete_droite, x_droite, HAUT_CONTENU, largeur, hauteur_entete)
    for forme in (entete_gauche, entete_droite):
        for rpr in forme.iter(qn("a:defRPr")):
            rpr.set("b", "1")
    contenu_gauche = zone(layout, 12)
    contenu_droite = zone(layout, 13)
    retyper(contenu_gauche, None, 2, "half")
    retyper(contenu_droite, None, 4, "quarter")
    placer(contenu_gauche, GAUCHE, haut, largeur, hauteur)
    placer(contenu_droite, x_droite, haut, largeur, hauteur)


def vers_content_with_caption(layout: SlideLayout) -> None:
    """« Page avec 1 photo + texte » : texte en haut (11), objet dessous (12).

    Pandoc emploie cette disposition pour un texte suivi d'une image ou d'un
    tableau. La photo étroite du modèle (un tiers de la largeur) écrasait les
    tableaux : texte et objet prennent ici toute la largeur, l'un sous l'autre.
    """
    titre_diapositive(layout)
    hauteur_texte = 1100000
    objet = zone(layout, 12)
    texte = zone(layout, 11)
    retyper(objet, None, 1)
    retyper(texte, "body", 2, "half")
    placer(texte, GAUCHE, HAUT_CONTENU, LARGEUR_UTILE, hauteur_texte)
    haut_objet = HAUT_CONTENU + hauteur_texte + 100000
    placer(objet, GAUCHE, haut_objet, LARGEUR_UTILE, BAS_CONTENU - haut_objet)


def vers_blank(layout: SlideLayout) -> None:
    """« Derniere page » : conservée telle quelle, son titre fixe n'est pas rempli."""


CORRESPONDANCES: dict[str, tuple[str, Callable[[SlideLayout], None]]] = {
    "Slide couv": ("Title Slide", vers_title_slide),
    "Slide texte seulement": ("Title and Content", vers_title_and_content),
    "Slide fond dégradé": ("Section Header", vers_section_header),
    "Slide texte 2 colonnes": ("Two Content", vers_two_content),
    "Slide texte + graphiques": ("Comparison", vers_comparison),
    "Page avec 1 photo + texte": ("Content with Caption", vers_content_with_caption),
    "Derniere page": ("Blank", vers_blank),
}
"""Disposition du modèle -> (nom attendu par Pandoc, transformation)."""


def construire(source: Path, sortie: Path) -> None:
    """Produit le fichier de référence à partir du modèle Cerema.

    Args:
        source: Modèle .potx du Cerema.
        sortie: Fichier .pptx à écrire.
    """
    presentation = lire_modele(source)
    supprimer_diapositives(presentation)
    for nom_modele, (nom_pandoc, transformer) in CORRESPONDANCES.items():
        layout = disposition(presentation, nom_modele)
        transformer(layout)
        layout._element.find(qn("p:cSld")).set("name", nom_pandoc)
    presentation.save(sortie)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    construire(Path(sys.argv[1]), Path(sys.argv[2]))
